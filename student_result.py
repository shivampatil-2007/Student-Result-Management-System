import tkinter as tk
from tkinter import messagebox, ttk
from openpyxl import Workbook, load_workbook
import os

FILE = "student_results.xlsx"


# Create Excel file if it does not exist
def create_file():
    if not os.path.exists(FILE):
        wb = Workbook()
        ws = wb.active
        ws.append([
            "Name", "Roll No.", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total Marks", "Percentage", "Result"
        ])
        wb.save(FILE)


# Clear input boxes
def clear_fields():
    for entry in entries:
        entry.delete(0, tk.END)


# Add Student
def add_student():
    try:
        name = name_entry.get()
        roll = roll_entry.get()
        student_class = class_entry.get()

        marks = [
            int(sub1_entry.get()),
            int(sub2_entry.get()),
            int(sub3_entry.get()),
            int(sub4_entry.get()),
            int(sub5_entry.get())
        ]

        if name == "" or roll == "" or student_class == "":
            messagebox.showerror("Error", "Please enter all details")
            return

        # Check marks
        for mark in marks:
            if mark < 0 or mark > 100:
                messagebox.showerror(
                    "Error", "Marks should be between 0 and 100"
                )
                return

        total = sum(marks)
        percentage = total / 5

        # Pass if every subject is 35 or more
        if all(mark >= 35 for mark in marks):
            result = "Pass"
        else:
            result = "Fail"

        create_file()

        wb = load_workbook(FILE)
        ws = wb.active

        ws.append([
            name, roll, student_class,
            marks[0], marks[1], marks[2],
            marks[3], marks[4],
            total, f"{percentage:.1f}%", result
        ])

        wb.save(FILE)

        messagebox.showinfo("Success", "Student record saved!")
        clear_fields()

    except ValueError:
        messagebox.showerror(
            "Error", "Please enter valid marks and Roll No."
        )


# Get Result
def get_result():
    roll = result_roll_entry.get()

    if roll == "":
        messagebox.showerror("Error", "Enter Roll No.")
        return

    create_file()

    wb = load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):

        if str(row[1]) == roll:
            result_text = (
                f"Name: {row[0]}\n"
                f"Roll No.: {row[1]}\n"
                f"Class: {row[2]}\n"
                f"Total Marks: {row[8]}\n"
                f"Percentage: {row[9]}\n"
                f"Result: {row[10]}"
            )

            messagebox.showinfo("Student Result", result_text)
            return

    messagebox.showerror("Not Found", "Student record not found.")


# Show All Results
def show_all():
    create_file()

    wb = load_workbook(FILE)
    ws = wb.active

    # Clear old data
    for item in tree.get_children():
        tree.delete(item)

    for row in ws.iter_rows(min_row=2, values_only=True):

        tree.insert("", tk.END, values=(
            row[0],
            row[1],
            row[2],
            row[8],
            row[9],
            row[10]
        ))


# ---------------- MAIN WINDOW ----------------

root = tk.Tk()
root.title("Student Result Management System")
root.geometry("850x650")

title = tk.Label(
    root,
    text="STUDENT RESULT MANAGEMENT",
    font=("Arial", 20, "bold")
)
title.pack(pady=15)


# ---------------- ADD STUDENT ----------------

add_frame = tk.LabelFrame(
    root,
    text="Add Student",
    font=("Arial", 12, "bold"),
    padx=10,
    pady=10
)
add_frame.pack(padx=20, pady=5, fill="x")


tk.Label(add_frame, text="Name").grid(row=0, column=0, padx=5, pady=5)
name_entry = tk.Entry(add_frame)
name_entry.grid(row=0, column=1, padx=5, pady=5)

tk.Label(add_frame, text="Roll No.").grid(row=0, column=2, padx=5, pady=5)
roll_entry = tk.Entry(add_frame)
roll_entry.grid(row=0, column=3, padx=5, pady=5)

tk.Label(add_frame, text="Class").grid(row=0, column=4, padx=5, pady=5)
class_entry = tk.Entry(add_frame)
class_entry.grid(row=0, column=5, padx=5, pady=5)


tk.Label(add_frame, text="Subject 1").grid(row=1, column=0, padx=5, pady=5)
sub1_entry = tk.Entry(add_frame, width=8)
sub1_entry.grid(row=1, column=1, padx=5, pady=5)

tk.Label(add_frame, text="Subject 2").grid(row=1, column=2, padx=5, pady=5)
sub2_entry = tk.Entry(add_frame, width=8)
sub2_entry.grid(row=1, column=3, padx=5, pady=5)

tk.Label(add_frame, text="Subject 3").grid(row=1, column=4, padx=5, pady=5)
sub3_entry = tk.Entry(add_frame, width=8)
sub3_entry.grid(row=1, column=5, padx=5, pady=5)


tk.Label(add_frame, text="Subject 4").grid(row=2, column=0, padx=5, pady=5)
sub4_entry = tk.Entry(add_frame, width=8)
sub4_entry.grid(row=2, column=1, padx=5, pady=5)

tk.Label(add_frame, text="Subject 5").grid(row=2, column=2, padx=5, pady=5)
sub5_entry = tk.Entry(add_frame, width=8)
sub5_entry.grid(row=2, column=3, padx=5, pady=5)


save_button = tk.Button(
    add_frame,
    text="💾 Save",
    command=add_student,
    width=15
)
save_button.grid(row=2, column=4, columnspan=2, pady=10)


entries = [
    name_entry,
    roll_entry,
    class_entry,
    sub1_entry,
    sub2_entry,
    sub3_entry,
    sub4_entry,
    sub5_entry
]


# ---------------- GET RESULT ----------------

result_frame = tk.LabelFrame(
    root,
    text="Get Result",
    font=("Arial", 12, "bold"),
    padx=10,
    pady=10
)
result_frame.pack(padx=20, pady=10, fill="x")


tk.Label(
    result_frame,
    text="Enter Roll No.:"
).pack(side="left", padx=10)

result_roll_entry = tk.Entry(result_frame)
result_roll_entry.pack(side="left", padx=10)

get_button = tk.Button(
    result_frame,
    text="🔍 Get Result",
    command=get_result
)
get_button.pack(side="left", padx=10)


# ---------------- SHOW ALL RESULTS ----------------

show_button = tk.Button(
    root,
    text="📋 Show All Results",
    command=show_all,
    width=20
)
show_button.pack(pady=10)


# Treeview
columns = (
    "Name",
    "Roll No.",
    "Class",
    "Total Marks",
    "Percentage",
    "Result"
)

tree = ttk.Treeview(
    root,
    columns=columns,
    show="headings",
    height=10
)

for column in columns:
    tree.heading(column, text=column)
    tree.column(column, width=120)

tree.pack(padx=20, pady=10, fill="both", expand=True)


# Exit button
exit_button = tk.Button(
    root,
    text="Exit",
    command=root.destroy,
    width=15
)
exit_button.pack(pady=10)


# Start program
create_file()
root.mainloop()
